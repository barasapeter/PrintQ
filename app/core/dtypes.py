from pathlib import Path
import mimetypes
import zipfile
import xml.etree.ElementTree as ET

from openpyxl import load_workbook
from pptx import Presentation

import magic
from pypdf import PdfReader
from PIL import Image


MIME_MAP = {
    "application/pdf": "pdf",
    "application/msword": "doc",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "application/vnd.ms-powerpoint": "ppt",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": "pptx",
    "application/vnd.ms-excel": "xls",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
    "application/vnd.oasis.opendocument.text": "odt",
    "application/vnd.oasis.opendocument.presentation": "odp",
    "application/vnd.oasis.opendocument.spreadsheet": "ods",
    "application/rtf": "rtf",
    "text/rtf": "rtf",
    "text/plain": "txt",
    "text/csv": "csv",
    "text/html": "html",
    "application/xml": "xml",
    "text/xml": "xml",
    "image/jpeg": "jpg",
    "image/png": "png",
    "image/gif": "gif",
    "image/webp": "webp",
    "image/tiff": "tiff",
    "image/bmp": "bmp",
    "image/svg+xml": "svg",
    "image/heic": "heic",
    "image/heif": "heif",
    "image/x-icon": "ico",
    "image/vnd.microsoft.icon": "ico",
}


def _resolve_file_type(filepath: str) -> str | None:
    mime_type = magic.from_file(filepath, mime=True)
    if not mime_type:
        mime_type, _ = mimetypes.guess_type(filepath)
    return MIME_MAP.get(mime_type)


class NoPrintableContentError(ValueError):
    pass


def _validate_text(filepath: str) -> None:
    try:
        Path(filepath).read_text(errors="ignore")
    except Exception as e:
        raise NoPrintableContentError(f"Unable to read text document: {e}") from e


def _validate_docx(filepath: str) -> None:
    try:
        with zipfile.ZipFile(filepath) as archive:
            archive.read("word/document.xml")
    except Exception as e:
        raise NoPrintableContentError(f"Invalid or corrupted DOCX document: {e}") from e


def _validate_xlsx(filepath: str) -> None:
    try:
        load_workbook(filepath, read_only=True, data_only=True)
    except Exception as e:
        raise NoPrintableContentError(f"Invalid or corrupted XLSX document: {e}") from e


def _validate_pptx(filepath: str) -> None:
    try:
        Presentation(filepath)
    except Exception as e:
        raise NoPrintableContentError(f"Invalid or corrupted PPTX document: {e}") from e


def _validate_pdf(filepath: str) -> None:
    try:
        PdfReader(filepath)
    except Exception as e:
        raise NoPrintableContentError(f"Invalid or corrupted PDF document: {e}") from e


def _validate_image(filepath: str) -> None:
    try:
        with Image.open(filepath) as img:
            img.verify()
        with Image.open(filepath) as img:
            img.load()
    except Exception as e:
        raise NoPrintableContentError(f"Invalid or corrupted image file: {e}") from e


_VALIDATORS = {
    "txt": _validate_text,
    "csv": _validate_text,
    "html": _validate_text,
    "xml": _validate_text,
    "rtf": _validate_text,
    "docx": _validate_docx,
    "xlsx": _validate_xlsx,
    "pptx": _validate_pptx,
    "pdf": _validate_pdf,
    "jpg": _validate_image,
    "png": _validate_image,
    "gif": _validate_image,
    "webp": _validate_image,
    "tiff": _validate_image,
    "bmp": _validate_image,
    "svg": _validate_image,
    "heic": _validate_image,
    "heif": _validate_image,
    "ico": _validate_image,
}


def validate_document(filepath: str) -> str:
    file_type = _resolve_file_type(filepath)

    if not file_type:
        raise NoPrintableContentError("Unsupported or unrecognized file type")

    validator = _VALIDATORS.get(file_type)

    if not validator:
        raise NoPrintableContentError(f"Validation not implemented for '{file_type}'")

    validator(filepath)
    return file_type


if __name__ == "__main__":
    test_files = [
        "Dockerfile",
        "empty.txt",
        "image.jpg",
        "document.pdf",
    ]
    for test_file in test_files:
        try:
            print(f"{test_file}: {validate_document(test_file)}")
        except NoPrintableContentError as e:
            print(f"{test_file}: ERROR - {e}")
